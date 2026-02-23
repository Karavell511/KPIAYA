from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from app.models import KPIInstance


def render_kpi_pdf(instance: KPIInstance) -> bytes:
    buf = BytesIO()
    pdf = canvas.Canvas(buf, pagesize=A4)
    y = 800
    pdf.drawString(50, y, f"KPI Report: {instance.id}")
    y -= 20
    pdf.drawString(50, y, f"Employee ID: {instance.employee_id}")
    y -= 20
    pdf.drawString(50, y, f"Period: {instance.period_start} - {instance.period_end} ({instance.month})")
    y -= 20
    pdf.drawString(50, y, f"Status: {instance.status.value}")
    y -= 30
    for metric in instance.metrics:
        pdf.drawString(50, y, f"Metric #{metric.template_metric_id}: fact={metric.fact_value}, result={metric.metric_result}, overtime={metric.overtime_result}")
        y -= 20
    pdf.drawString(50, y, f"Overtime Total: {instance.overtime_amount}")
    y -= 20
    pdf.drawString(50, y, f"TOTAL: {instance.total_amount}")
    pdf.showPage()
    pdf.save()
    return buf.getvalue()


def render_kpi_xlsx(instance: KPIInstance) -> bytes:
    wb = Workbook()
    ws_summary = wb.active
    ws_summary.title = "Summary"
    ws_summary.append(["KPI ID", instance.id])
    ws_summary.append(["Employee ID", instance.employee_id])
    ws_summary.append(["Month", instance.month])
    ws_summary.append(["Period Start", str(instance.period_start)])
    ws_summary.append(["Period End", str(instance.period_end)])
    ws_summary.append(["Status", instance.status.value])
    ws_summary.append(["Total", float(instance.total_amount)])

    ws_metrics = wb.create_sheet("Metrics")
    ws_metrics.append(["Template Metric ID", "Fact Value", "Metric Result"])
    for metric in instance.metrics:
        ws_metrics.append([metric.template_metric_id, float(metric.fact_value), float(metric.metric_result)])

    ws_ot = wb.create_sheet("Overtime")
    ws_ot.append(["Template Metric ID", "Overtime"])
    for metric in instance.metrics:
        ws_ot.append([metric.template_metric_id, float(metric.overtime_result or 0)])
    ws_ot.append(["Total Overtime", float(instance.overtime_amount or 0)])

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def render_zip(instances: list[KPIInstance], fmt: str) -> bytes:
    buf = BytesIO()
    with ZipFile(buf, "w", compression=ZIP_DEFLATED) as archive:
        for instance in instances:
            if fmt == "pdf":
                archive.writestr(f"kpi_{instance.id}.pdf", render_kpi_pdf(instance))
            else:
                archive.writestr(f"kpi_{instance.id}.xlsx", render_kpi_xlsx(instance))
    return buf.getvalue()
